from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_ivy_word_report import (
    BLACK,
    BORDER,
    INK,
    MUTED,
    PRIMARY,
    PRIMARY_LIGHT,
    SAGE_LIGHT,
    WHITE,
    add_bullet,
    add_manual_numbered,
    add_page_number,
    add_paragraph,
    prevent_row_split,
    set_cell_margins,
    set_cell_width,
    set_font,
    set_repeat_table_header,
    set_table_borders,
    shade_cell,
    style_table_text,
)
from extension_assessment_content import (
    COMMON_RESULT_FLOW,
    DOMINANT_PATTERN_SCORING,
    EXTENSION_ASSESSMENTS,
)


OUTPUT = Path(__file__).resolve().parent / "outputs" / "好理家在財務探索_1加4檢測架構討論稿.docx"


TYPE_DEFINITIONS = [
    {
        "code": "安全規劃型",
        "public": "重視安心與準備",
        "meaning": "做決定前，會先確認基本生活與預留空間。",
        "next": "我的財務計畫卡在哪裡？",
        "question": "計畫遇到變化時，你還能繼續做下去嗎？",
    },
    {
        "code": "當下滿足型",
        "public": "重視當下需要",
        "meaning": "會把眼前生活感受與現在真正需要的事情放進考量。",
        "next": "壓力來時，我會怎麼用錢？",
        "question": "哪些花費真的在照顧自己，哪些比較像短暫紓壓？",
    },
    {
        "code": "精算判斷型",
        "public": "先比較再決定",
        "meaning": "會先看清金額、條件、順序與後續影響。",
        "next": "我的財務計畫卡在哪裡？",
        "question": "想得很清楚之後，計畫通常走得下去嗎？",
    },
    {
        "code": "行動嘗試型",
        "public": "願意先試再調整",
        "meaning": "願意先做一個負擔得起的小嘗試，再依結果調整。",
        "next": "我看得懂自己的錢嗎？",
        "question": "開始以前，你掌握可用金額與後續負擔嗎？",
    },
    {
        "code": "關係照顧型",
        "public": "會先顧到重要的人",
        "meaning": "會把家人、重要的人與共同生活放進決定。",
        "next": "我和家人的金錢界線清楚嗎？",
        "question": "照顧重要的人時，你也清楚自己能負擔到哪裡嗎？",
    },
    {
        "code": "成就證明型",
        "public": "希望努力有看得見的成果",
        "meaning": "希望金錢選擇能支持目標，也讓努力有被看見的成果。",
        "next": "壓力來時，我會怎麼用錢？",
        "question": "想讓努力有成果時，花錢會不會也成為獎勵或壓力出口？",
    },
    {
        "code": "彈性應變型",
        "public": "先保留調整空間",
        "meaning": "會保留選擇空間，遇到變化時再重新安排。",
        "next": "我看得懂自己的錢嗎？",
        "question": "調整安排後，你仍知道錢去了哪裡、還剩多少嗎？",
    },
    {
        "code": "延後面對型",
        "public": "會先留點空間再處理",
        "meaning": "壓力變大時，可能先暫緩，等比較能面對時再回來處理。",
        "next": "我看得懂自己的錢嗎？",
        "question": "如果只先看一小部分，你最願意從哪項資訊開始？",
    },
]


ENTRY_QUESTIONS = [
    {
        "title": "多出一筆可以安排的錢",
        "scenario": "這個月多出一筆可以自己安排的 3,000 元。你最可能先怎麼做？",
        "options": [
            ("A", "先留起來，讓生活臨時有需要時還有空間。", "安全規劃型"),
            ("B", "拿一部分改善現在的生活，或做一件期待已久的事。", "當下滿足型"),
            ("C", "先想到家人或重要的人最近有沒有共同需要。", "關係照顧型"),
            ("D", "用在一個想完成的目標，讓最近的努力有明確成果。", "成就證明型"),
        ],
    },
    {
        "title": "需要馬上決定的機會",
        "scenario": "你看到一門可能對工作或生活有幫助的課程，但優惠今晚截止。你最可能怎麼決定？",
        "options": [
            ("A", "先看它能不能讓重要目標往前走，再決定是否投入。", "成就證明型"),
            ("B", "先從短期或入門方案試試看，不一次投入太多。", "行動嘗試型"),
            ("C", "比較內容、總費用與其他選擇，再做決定。", "精算判斷型"),
            ("D", "選擇可以退改或調整的方案，替之後保留空間。", "彈性應變型"),
        ],
    },
    {
        "title": "重要的人臨時需要幫忙",
        "scenario": "家人或重要的人臨時開口，希望你幫忙一筆超出平常往來的錢。你最可能先怎麼做？",
        "options": [
            ("A", "先了解對方真正需要什麼，再一起談可以怎麼處理。", "關係照顧型"),
            ("B", "先確認不影響自己的基本生活，再說能幫多少。", "安全規劃型"),
            ("C", "除了直接給錢，也一起找分期、資源或其他做法。", "彈性應變型"),
            ("D", "先請對方給我一點時間，等自己比較能面對時再回覆。", "延後面對型"),
        ],
    },
    {
        "title": "突然出現的費用",
        "scenario": "家裡的電器、交通工具或醫療突然需要一筆原本沒有安排的費用。你最可能先怎麼處理？",
        "options": [
            ("A", "先處理最急的一小部分，讓事情不要繼續擴大。", "行動嘗試型"),
            ("B", "先確認金額、期限與不同做法的總負擔。", "精算判斷型"),
            ("C", "先讓自己緩一下，等比較有空間時再回來整理。", "延後面對型"),
            ("D", "先把必要生活費留好，再調整其他安排。", "安全規劃型"),
        ],
    },
    {
        "title": "想買東西獎勵自己",
        "scenario": "完成一件辛苦很久的事情後，你想買一樣不是現在一定需要的東西。你最可能怎麼想？",
        "options": [
            ("A", "先比較價格、使用頻率與買下後對其他支出的影響。", "精算判斷型"),
            ("B", "如果它能代表這段努力的成果，我會先訂預算再選擇。", "成就證明型"),
            ("C", "如果它能讓現在的生活舒服一點，我會在負擔得起時買下。", "當下滿足型"),
            ("D", "先用租借、試用或小額版本，確認適合再投入。", "行動嘗試型"),
        ],
    },
    {
        "title": "需要一起出錢的計畫",
        "scenario": "你和家人或重要的人正討論一個需要一起出錢的計畫，例如旅行、搬家或添購家用品。你最可能先在意什麼？",
        "options": [
            ("A", "這個安排能不能讓大家現在的生活多一點好的感受。", "當下滿足型"),
            ("B", "預算能不能保留調整空間，遇到變化時還改得動。", "彈性應變型"),
            ("C", "這個計畫能不能成為大家一起完成的重要里程碑。", "成就證明型"),
            ("D", "每個人的期待、負擔與可以接受的範圍有沒有說清楚。", "關係照顧型"),
        ],
    },
    {
        "title": "好幾件要用錢的事一起出現",
        "scenario": "帳單、生活安排和一個想完成的目標同時需要你處理。你最可能從哪裡開始？",
        "options": [
            ("A", "先看哪些事情能移動，再依當下狀況重新排序。", "彈性應變型"),
            ("B", "先放一下，等壓力沒有那麼大時再回來看。", "延後面對型"),
            ("C", "先完成一件今天做得到的小事，讓整理開始動起來。", "行動嘗試型"),
            ("D", "先把金額、期限與優先順序列清楚。", "精算判斷型"),
        ],
    },
    {
        "title": "接下來的收入可能變少",
        "scenario": "你得知接下來一段時間的收入可能變少，原本的支出安排需要調整。你最可能先怎麼做？",
        "options": [
            ("A", "先不急著看全部，等自己比較能面對時再整理。", "延後面對型"),
            ("B", "先和可能受到影響的人說明，再一起調整安排。", "關係照顧型"),
            ("C", "先保留基本生活與不能延後的支出。", "安全規劃型"),
            ("D", "即使要縮減，也保留一項能維持生活感受的小支出。", "當下滿足型"),
        ],
    },
]


EXTENSIONS = [
    {
        "title": "我看得懂自己的錢嗎",
        "purpose": "整理哪些金錢資訊已經掌握，哪些還需要確認。",
        "topics": [
            "一般月份大約有多少收入可以使用",
            "每月一定要負擔的生活項目與金額",
            "債務餘額、每期金額與繳款時間；沒有債務者使用獨立選項",
            "扣除基本生活與還款後，可以調整的金額",
            "未來一至三個月可能出現的較大支出",
            "不確定時，是否知道去哪裡找紀錄並重新整理",
        ],
        "results": "掌握得較清楚／已有大致輪廓／邊走邊確認／適合先從一項開始",
        "available": "財務月報表、記帳助理、財務健康儀表板",
        "planned": "收支盤點引導",
    },
    {
        "title": "壓力來時我會怎麼用錢",
        "purpose": "看見哪些壓力會改變用錢行為，以及自己通常如何回應。",
        "topics": [
            "疲累或受挫後，第一個用錢反應",
            "情緒較強時，遇到限時商品會如何決定",
            "收到帳單或看到餘額下降時的第一個反應",
            "覺得虧欠自己或家人時，是否會用花費補償",
            "發生臨時支出後，會彈性調整還是全面收緊",
            "發現衝動消費或暫緩處理後，下一步會怎麼做",
        ],
        "results": "停一下再決定／先讓自己好過一點／暫時先不碰／先把支出收緊／用花費照顧自己或他人",
        "available": "財務焦慮檢測、記帳助理、生命歷程整理",
        "planned": "生活財務覺察卡、用錢情境紀錄、情緒與消費內容",
    },
    {
        "title": "我的家庭金錢界線",
        "purpose": "整理關係中的期待、責任與決定權如何影響金錢選擇。",
        "topics": [
            "面對未預期的金錢請求時，能否先確認負擔範圍",
            "家庭固定支出如何分工，是否能重新討論",
            "被要求先刷卡、借名或代為承擔時會如何處理",
            "是否能向家人說明自己的需要與限制",
            "家中較大的支出通常由誰決定",
            "發現分工長期不均時，是否有可行的調整方式",
        ],
        "results": "比較能一起協商／責任較多落在你身上／常把自己放在後面／多半各自處理／決定權較集中",
        "available": "財務合拍檢測、財務月報表、生命歷程整理",
        "planned": "家庭金錢對話工具、家庭資源與責任圖、談錢不傷感情課程",
    },
    {
        "title": "我的財務計畫卡在哪裡",
        "purpose": "整理知道要做什麼以後，最容易在哪個環節停住。",
        "topics": [
            "能否把想改善的事情說成一個明確目標",
            "能否把目標拆成短時間內可完成的行動",
            "是否有簡單而能持續的追蹤方式",
            "發生臨時變化時，能否縮小或調整原定行動",
            "中斷一次後，是否知道如何重新開始",
            "目前是因為目標太多、標準太高或生活責任而停住",
        ],
        "results": "穩定推進／同時想做很多／起步快但較難維持／想準備好再開始／常被生活變化打斷",
        "available": "生活目標財務規劃、儲蓄目標試算、時間資源盤點助理",
        "planned": "30 天財務小行動、馴錢養成班、行動追蹤工具",
    },
]


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), PRIMARY)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    title_style = styles["Title"]
    title_style.font.name = "Microsoft JhengHei"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    title_style.paragraph_format.space_after = Pt(11)
    title_ppr = title_style._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)

    for style_name, size, before, after in (
        ("Heading 1", 16.5, 17, 7),
        ("Heading 2", 13.5, 12, 5),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = styles[style_name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[list_style_name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.2

    header_paragraph = section.header.paragraphs[0]
    header_run = header_paragraph.add_run("好理家在  財務探索檢測討論稿")
    set_font(header_run, size=8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    properties = doc.core_properties
    properties.title = "好理家在財務探索檢測設計討論稿"
    properties.subject = "一個入口檢測與四個延伸主題"
    properties.author = "好理家在"
    properties.keywords = "財務探索, 入口檢測, 延伸檢測, 使用者旅程"


def keep_table_together(table):
    """Keep a short table on one page when it fits."""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def add_data_table(
    doc,
    headers,
    rows,
    widths,
    font_size=9.2,
    center_columns=None,
    compact=False,
    keep_together=False,
):
    center_columns = set(center_columns or [])
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, (label, width) in enumerate(zip(headers, widths)):
        cell = header.cells[index]
        set_cell_width(cell, width)
        vertical_margin = 88 if compact else 120
        set_cell_margins(
            cell,
            top=vertical_margin,
            start=120,
            bottom=vertical_margin,
            end=120,
        )
        shade_cell(cell, PRIMARY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        style_table_text(
            cell,
            size=9.3,
            color=WHITE,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER if index in center_columns else WD_ALIGN_PARAGRAPH.LEFT,
        )
        cell.paragraphs[0].text = label
        style_table_text(
            cell,
            size=9.3,
            color=WHITE,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER if index in center_columns else WD_ALIGN_PARAGRAPH.LEFT,
        )

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for column_index, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[column_index]
            set_cell_width(cell, width)
            vertical_margin = 72 if compact else 105
            set_cell_margins(
                cell,
                top=vertical_margin,
                start=120,
                bottom=vertical_margin,
                end=120,
            )
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                shade_cell(cell, SAGE_LIGHT)
            cell.text = str(value)
            style_table_text(
                cell,
                size=font_size,
                color=INK,
                bold=column_index == 0,
                align=WD_ALIGN_PARAGRAPH.CENTER if column_index in center_columns else WD_ALIGN_PARAGRAPH.LEFT,
            )
    if keep_together:
        keep_table_together(table)
    spacing = doc.add_paragraph()
    spacing.paragraph_format.space_before = Pt(0)
    spacing.paragraph_format.space_after = Pt(2)
    return table


def add_architecture_overview(doc):
    entrance = doc.add_table(rows=1, cols=1)
    entrance.alignment = WD_TABLE_ALIGNMENT.CENTER
    entrance.autofit = False
    set_table_borders(entrance)
    cell = entrance.cell(0, 0)
    set_cell_width(cell, Cm(16.7))
    set_cell_margins(cell, top=170, start=180, bottom=170, end=180)
    shade_cell(cell, PRIMARY_LIGHT)
    cell.text = "入口檢測  做金錢決定時我通常先在意什麼  8 題約 2 分鐘\n輸出  這次較常出現的一至兩種考量"
    style_table_text(cell, size=11, color=INK, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    arrow = doc.add_paragraph()
    arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    arrow.paragraph_format.space_before = Pt(3)
    arrow.paragraph_format.space_after = Pt(3)
    arrow_run = arrow.add_run("↓  依結果優先推薦  也可以自己選擇想了解的主題")
    set_font(arrow_run, size=9.5, color=MUTED)

    grid = doc.add_table(rows=2, cols=2)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    set_table_borders(grid)
    labels = [
        "我看得懂自己的錢嗎\n6 題  整理財務掌握",
        "壓力來時我會怎麼用錢\n6 題  整理壓力反應",
        "我的家庭金錢界線\n6 題  整理關係與責任",
        "我的財務計畫卡在哪裡\n6 題  整理執行阻力",
    ]
    for index, cell in enumerate([grid.cell(0, 0), grid.cell(0, 1), grid.cell(1, 0), grid.cell(1, 1)]):
        set_cell_width(cell, Cm(8.35))
        set_cell_margins(cell, top=150, start=150, bottom=150, end=150)
        shade_cell(cell, SAGE_LIGHT if index % 2 == 0 else "F5F0EC")
        cell.text = labels[index]
        style_table_text(cell, size=10.2, color=INK, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    second_arrow = doc.add_paragraph()
    second_arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    second_arrow.paragraph_format.space_before = Pt(3)
    second_arrow.paragraph_format.space_after = Pt(3)
    second_arrow_run = second_arrow.add_run("↓  每完成一項  就加入一塊已確認的觀察")
    set_font(second_arrow_run, size=9.5, color=MUTED)

    puzzle = doc.add_table(rows=1, cols=5)
    puzzle.alignment = WD_TABLE_ALIGNMENT.CENTER
    puzzle.autofit = False
    set_table_borders(puzzle)
    pieces = ["用錢習慣", "財務掌握", "壓力反應", "家庭界線", "計畫執行"]
    for index, cell in enumerate(puzzle.rows[0].cells):
        set_cell_width(cell, Cm(3.34))
        set_cell_margins(cell, top=120, start=80, bottom=120, end=80)
        if index == 0:
            shade_cell(cell, PRIMARY_LIGHT)
        cell.text = pieces[index]
        style_table_text(cell, size=9, color=INK, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

def add_question(doc, number, item):
    heading = doc.add_heading(f"第 {number} 題  {item['title']}", level=3)
    if number in (1, 3, 5, 7):
        heading.paragraph_format.page_break_before = number != 1
    heading.paragraph_format.space_after = Pt(2)
    scenario = add_paragraph(doc, item["scenario"])
    scenario.paragraph_format.space_after = Pt(4)
    scenario.paragraph_format.keep_with_next = True
    rows = [(label, text, direction) for label, text, direction in item["options"]]
    rows.append(("E", "我沒有遇過這個狀況", "不計分"))
    add_data_table(
        doc,
        ["選項", "使用者看到的答案", "內部觀察"],
        rows,
        [Cm(1.3), Cm(12.0), Cm(3.4)],
        font_size=9.0,
        center_columns={0, 2},
        compact=True,
    )


def add_extension_question(doc, extension_number, question_number, item):
    heading = doc.add_heading(
        f"延伸 {extension_number}｜第 {question_number} 題  {item['title']}",
        level=3,
    )
    if question_number > 1:
        heading.paragraph_format.page_break_before = True
    heading.paragraph_format.space_after = Pt(2)

    scenario = add_paragraph(doc, f"情境  {item['scenario']}", bold_lead="情境")
    scenario.paragraph_format.space_after = Pt(2)
    scenario.paragraph_format.keep_with_next = True
    prompt = add_paragraph(doc, item["prompt"])
    prompt.paragraph_format.space_after = Pt(4)
    prompt.paragraph_format.keep_with_next = True

    rows = [
        (option["label"], option["text"], option["internal"])
        for option in item["options"]
    ]
    skip_label = chr(ord("A") + len(item["options"]))
    rows.append((skip_label, "我沒有遇過這個狀況", "完成但不計分"))
    add_data_table(
        doc,
        ["選項", "使用者看到的答案", "內部決策訊號"],
        rows,
        [Cm(1.25), Cm(11.65), Cm(3.8)],
        font_size=8.7,
        center_columns={0, 2},
        compact=True,
    )


def add_extension_result_spec(doc, item):
    result_heading = doc.add_heading("完成六題後｜決策覺察結果", level=2)
    result_heading.paragraph_format.page_break_before = True

    doc.add_heading("暫定計分", level=3)
    scoring_items = item["scoring"] or DOMINANT_PATTERN_SCORING
    for rule in scoring_items:
        add_bullet(doc, rule)

    doc.add_heading("可能呈現的決策方向", level=3)
    result_rows = [
        (result["name"], result["path"], result["step"])
        for result in item["results"]
    ]
    add_data_table(
        doc,
        ["本次決策方向", "從答案整理出的決策路徑", "先整理一小步"],
        result_rows,
        [Cm(3.3), Cm(8.45), Cm(4.95)],
        font_size=8.35,
        compact=True,
    )

    tool_heading = doc.add_heading("使用者主動展開後｜工具建議", level=3)
    tool_heading.paragraph_format.page_break_before = True
    tool_rows = [
        (result["name"], result["primary"], result["bridge"], result["others"])
        for result in item["results"]
    ]
    add_data_table(
        doc,
        ["本次決策方向", "主要工具", "承接剛才的小步驟", "其他資源"],
        tool_rows,
        [Cm(3.35), Cm(3.0), Cm(6.35), Cm(4.0)],
        font_size=8.2,
        compact=True,
    )
    add_paragraph(
        doc,
        f"若有效回答不足，先提供通用起點「{item['fallback_tool']}」，不再宣稱這是個人化推薦。",
    )
    if item.get("safety_note"):
        add_paragraph(
            doc,
            f"家庭安全邊界  {item['safety_note']}",
            bold_lead="家庭安全邊界",
        )


def add_extension_assessment(doc, number, item):
    heading = doc.add_heading(f"延伸檢測 {number}  {item['title']}", level=1)
    heading.paragraph_format.page_break_before = True
    add_paragraph(doc, f"目的  {item['purpose']}", bold_lead="目的")
    add_paragraph(doc, f"覺察重點  {item['awareness_focus']}", bold_lead="覺察重點")
    for question_number, question in enumerate(item["questions"], start=1):
        add_extension_question(doc, number, question_number, question)
    add_extension_result_spec(doc, item)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title_run = title.add_run("好理家在財務探索檢測設計討論稿")
    set_font(title_run, size=24, bold=True, color=BLACK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run("一個入口檢測與四個延伸主題  2026年9月18日")
    set_font(subtitle_run, size=11, color=MUTED)

    add_paragraph(
        doc,
        "本稿規劃一個 8 題入口與四個 6 題延伸檢測。入口整理做金錢決定時較常先在意什麼；延伸檢測補充決策情境、第一反應與可能影響。",
    )
    add_paragraph(
        doc,
        "目前定位  內容確認與試填草案；已列出 32 題、暫定計分、結果與工具路徑，仍需認知訪談及小型試填。",
        bold_lead="目前定位",
    )

    section_one = doc.add_heading("一  一個入口與四個延伸主題", level=1)
    section_one.paragraph_format.page_break_before = True
    add_architecture_overview(doc)
    add_paragraph(
        doc,
        "使用方式  入口先整理主要考量，再推薦一項延伸主題；延伸沒有固定順序，使用者可改選或停止。結果只描述本次決策線索，不作人格或能力判定。",
        bold_lead="使用方式",
    )

    module_heading = doc.add_heading("五項探索回答的問題", level=2)
    module_heading.paragraph_format.page_break_before = True
    module_heading.paragraph_format.keep_with_next = True
    module_rows = [
        ("入口探索", "我做金錢決定時通常先在意什麼", "一至兩個考量線索與反思問題"),
        ("財務掌握", "決定以前 我通常掌握多少資訊", "資訊如何進入決定與一個整理起點"),
        ("壓力與用錢", "壓力出現時 我先保護什麼", "觸發情境 第一反應與可能影響"),
        ("家庭金錢界線", "家人需要出現時 我怎麼衡量責任與界線", "決定方式 可能影響與一個小步驟"),
        ("計畫執行", "生活變動時 我在哪一步改變原本決定", "行動阻力 決策轉向與最低版本"),
    ]
    add_data_table(
        doc,
        ["探索主題", "回答的問題", "結果形式"],
        module_rows,
        [Cm(3.0), Cm(7.4), Cm(6.3)],
        font_size=9.1,
        compact=True,
        keep_together=True,
    )

    section_two = doc.add_heading("二  入口檢測的八種傾向", level=1)
    add_paragraph(
        doc,
        "內部保留工作名稱；使用者端採中性描述，不顯示人格標籤。",
    )
    type_rows = [
        (item["code"], item["public"], item["meaning"])
        for item in TYPE_DEFINITIONS
    ]
    add_data_table(
        doc,
        ["內部工作名稱", "建議對外名稱", "這個傾向在意什麼"],
        type_rows,
        [Cm(3.0), Cm(3.8), Cm(9.9)],
        font_size=8.9,
    )

    route_heading = doc.add_heading("入口結果與下一個主題", level=2)
    route_heading.paragraph_format.page_break_before = True
    route_heading.paragraph_format.keep_with_next = True
    route_rows = [
        (item["public"], item["question"], item["next"])
        for item in TYPE_DEFINITIONS
    ]
    add_data_table(
        doc,
        ["入口結果", "想知道更多", "優先推薦"],
        route_rows,
        [Cm(3.4), Cm(8.9), Cm(4.4)],
        font_size=8.8,
        compact=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        "原則  依主要考量推薦一項；只有兩項主要考量並列且指向不同主題時，才提供兩個等位選擇。",
        bold_lead="原則",
    )

    section_three = doc.add_heading("三  入口檢測 8 題草案", level=1)
    section_three.paragraph_format.page_break_before = True
    add_paragraph(
        doc,
        "檢測名稱  「做金錢決定時，我通常先在意什麼？」每題選一個最接近日常第一反應的做法；內部訊號不顯示，「我沒有遇過這個狀況」完成但不計分。",
        bold_lead="檢測名稱",
    )
    for number, item in enumerate(ENTRY_QUESTIONS, start=1):
        add_question(doc, number, item)

    doc.add_heading("四  入口計分與結果", level=1)
    doc.add_heading("入口計分", level=2)
    for item in [
        "每個答案只對應一項傾向並記一次，不使用隱藏加權。",
        "全部有效時比較被選次數；有「未遇過」時，改以被選次數除以有效出現機會。",
        "至少六題可計分；主要與次要考量都須實際被選至少兩次。線索不足時不排名。",
        "兩項最高並列就並列呈現；三項以上寫成多種考量並存，不使用隱藏規則破同分。",
        "結果不顯示分數、不與他人比較；相對選擇也不適合畫八軸雷達圖。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("結果頁內容順序", level=2)
    result_steps = [
        "顯示主要、次要或並列考量，不寫成固定人格。",
        "用實際答案整理「情境 → 先在意 → 第一反應 → 幫助與限制」。",
        "讓使用者確認是否貼近；確認不修改原始計分。",
        "提供一個短時間可完成的小步驟。",
        "推薦一項延伸主題；使用者可改選或停止。工具僅在延伸結果後主動展開。",
    ]
    for number, item in enumerate(result_steps, start=1):
        add_manual_numbered(doc, number, item)

    example_heading = doc.add_heading("結果文字範例", level=2)
    add_paragraph(
        doc,
        "這次較常出現的考量  重視關係需要  視情況調整",
        bold_lead="這次較常出現的考量",
    )
    add_paragraph(
        doc,
        "你做金錢決定時，常會先想到重要的人是否受到影響；情況改變時，也傾向重新找一個大家比較能接受的做法。這讓你比較能照顧共同生活，但也可能在調整別人的需要時，還沒先確認自己的負擔範圍。",
    )
    add_paragraph(
        doc,
        "下一步可以先做  回覆金錢請求以前，先寫下不影響基本生活的可幫範圍。",
        bold_lead="下一步可以先做",
    )
    add_paragraph(
        doc,
        "想知道更多  你能幫到哪裡，又如何替自己保留生活餘裕？可以接著探索「我的家庭金錢界線」，也可以先選其他主題。",
        bold_lead="想知道更多",
    )

    section_five = doc.add_heading("五  四個延伸檢測完整題目與結果", level=1)
    section_five.paragraph_format.page_break_before = True
    add_paragraph(
        doc,
        "內容  24 題生活情境選擇題、內部決策訊號、暫定結果與工具路徑。",
        bold_lead="內容",
    )
    add_paragraph(
        doc,
        "作答規則  每項 6 題；選答案後直接進下一題，未答不能前進。「我沒有遇過這個狀況」完成但不計分。作答頁不顯示分數、內部訊號或工具。",
        bold_lead="作答規則",
    )
    add_paragraph(
        doc,
        "結果與工具  小步驟先整理主工具需要的起始資料；再以「用剛剛整理的內容繼續」展開一項主要工具及最多兩項其他資源。現階段是內容銜接，不表示資料會自動帶入；自動串接、登入與保存需另行驗收。",
        bold_lead="結果與工具",
    )
    for number, item in enumerate(COMMON_RESULT_FLOW, start=1):
        add_manual_numbered(doc, number, item)
    for number, item in enumerate(EXTENSION_ASSESSMENTS, start=1):
        add_extension_assessment(doc, number, item)

    doc.add_heading("六  使用者旅程與互動", level=1)
    for item in [
        "進度顯示目前回答到第幾題，例如第 3／8 題，不以已選答案數量當作進度。",
        "滑鼠或觸控選到答案後，短暫保留選取狀態再直接進下一題；鍵盤操作保留下一題按鈕。",
        "未選答案不能前進；返回上一題時保留本次作答，卡片高度盡量固定，避免畫面上下跳動。",
        "答案只留在當次分頁記憶體；重新整理、關閉分頁或重新檢測都會清除，且不寫入會員紀錄。進入既有工具後是否另存資料，依該工具自己的規則說明。",
        "若日後要保存財務拼圖，必須由使用者主動選擇加入，並說明保存位置、期限及刪除方式。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("跨主題整理方式", level=2)
    add_paragraph(
        doc,
        "只串接已完成主題，不推測未作答內容、不加總分，也不稱為完整財務評估。",
    )

    doc.add_heading("七  製作與驗證重點", level=1)
    add_paragraph(
        doc,
        "實作順序  先推出入口檢測；延伸檢測依試填結果與使用者選擇逐項上線。未上線項目只蒐集探索意願，不提供空白入口。",
        bold_lead="實作順序",
    )
    doc.add_heading("正式製作前的驗證", level=2)
    for item in [
        "內容負責人確認八種傾向的定義、差異與題目映射；無法對齊時改稱考量線索。",
        "認知訪談確認題意、選項共鳴與缺少的反應。",
        "行動版檢查點選、返回、進度、卡片穩定與結果理解。",
        "小型試填觀察完成時間、離開率、選項分布、同分比例與結果共鳴。",
        "使用者看完延伸結果後，能說出自己的情境、優先考量與第一反應。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("先不做", level=2)
    for item in [
        "未驗證前一次上線四項延伸檢測。",
        "使用入口八軸雷達圖或跨人比較。",
        "未經同意保存跨次作答。",
        "在結果與小步驟前顯示工具，或把探索結果稱為完整評估。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("八  後續確認事項", level=1)
    decisions = [
        "入口八種傾向的名稱、定義與單一答案映射是否成立。",
        "「未遇過」、最低有效題數與並列規則是否採用本稿方案。",
        "入口一般推薦一項；兩項主要考量並列時才提供兩個等位選擇，是否確認。",
        "四項延伸的 24 題與決策覺察路徑是否符合預期。",
        "延伸至少五題有效才產生主要結果；同分不強制排序，是否確認。",
        "工具是否維持結果後主動展開，一項主要工具加最多兩項其他資源。",
        "作答是否維持當次分頁暫存；跨次保存需另訂同意與刪除規則。",
        "是否先製作入口 MVP，再依試填結果決定首項延伸檢測。",
    ]
    for number, item in enumerate(decisions, start=1):
        add_manual_numbered(doc, number, item)

    reference_heading = doc.add_heading("參考資料", level=1)
    reference_heading.paragraph_format.page_break_before = True
    source_one = doc.add_paragraph()
    source_one.add_run("好理家在財務工具  ")
    add_hyperlink(source_one, "https://www.familyfinhealth.com/toolbox", "https://www.familyfinhealth.com/toolbox")
    source_two = doc.add_paragraph()
    source_two.add_run("好理家在開始檢測  ")
    add_hyperlink(source_two, "https://www.familyfinhealth.com/assessment", "https://www.familyfinhealth.com/assessment")
    source_three = doc.add_paragraph()
    source_three.add_run("問卷認知訪談方法  ")
    add_hyperlink(
        source_three,
        "CDC Cognitive Interviewing",
        "https://www.cdc.gov/nchs/ccqder/question-evaluation/cognitive-interviewing.html",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
