from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parent / "outputs" / "做金錢決定時我會在意什麼_檢測設計討論稿.docx"

ANSWER_OPTIONS = [
    "很不像我",
    "比較不像我",
    "比較像我",
    "很像我",
    "我沒有遇過這個狀況",
]

BLACK = "000000"
INK = "4E403A"
MUTED = "77655D"
PRIMARY = "8F4D61"
PRIMARY_LIGHT = "F7EDEF"
SAGE_LIGHT = "EEF2EC"
BLUE_LIGHT = "EEF3F7"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


SCENARIOS = [
    {
        "title": "情境一  多出一筆可以安排的錢",
        "situation": "這個月手上多出一筆可以自己安排的錢，例如獎金、退費或臨時多出的收入。",
        "purpose": "觀察面對可運用的資源時，會先保留安全感、安排用途，還是照顧現在的生活。",
        "questions": [
            (1, "我會先留下一部分，生活臨時有需要時可以使用。", "安心守護"),
            (2, "我會先分配這筆錢要用在哪些地方，再開始使用。", "精算規劃"),
            (3, "我會想拿一部分改善現在的生活，或完成一直想做的事。", "當下感受"),
        ],
    },
    {
        "title": "情境二  重要的人臨時需要幫忙",
        "situation": "家人或重要的人臨時開口，希望你能幫忙一筆錢。",
        "purpose": "觀察自己的生活安全、彼此關係及當下心力，會如何一起影響金錢決定。",
        "questions": [
            (4, "我會先聽對方說明，也讓對方知道我的顧慮，再一起想怎麼處理。", "關係照顧"),
            (5, "一想到要做決定，我可能會先放著，等自己比較能面對時再回來處理。", "壓力迴避"),
            (6, "我會先確認幫忙後，自己的基本生活還過得去，再決定能幫多少。", "安心守護"),
        ],
    },
    {
        "title": "情境三  很快就要決定的機會",
        "situation": "有一個你很心動、需要花一筆錢，而且很快就要決定的機會。",
        "purpose": "觀察面對有時間限制的選擇時，會優先考慮眼前改善、先試著行動，還是完整計算後再決定。",
        "questions": [
            (7, "它能不能立即改善生活，或帶來我想要的體驗，會影響我的決定。", "當下感受"),
            (8, "即使還不確定最後值不值得，我也願意先投入一小部分試試看。", "機會行動"),
            (9, "我會先算清楚總共要花多少、之後還會不會有其他費用，再決定。", "精算規劃"),
        ],
    },
    {
        "title": "情境四  突然出現的費用",
        "situation": "家裡的電器、交通工具或醫療突然需要一筆原本沒有安排的費用。",
        "purpose": "觀察遇到意外壓力時，會先緩一緩、守住基本生活，還是尋找其他解決方式。",
        "questions": [
            (10, "壓力一下子上來時，我可能會先不看金額或處理方式，等自己比較能面對時再處理。", "壓力迴避"),
            (11, "我會先把必要的生活開銷留好，再想其他地方怎麼調整。", "安心守護"),
            (12, "我會先找一個現在做得到的方法，例如詢問可用資源或改用替代做法。", "機會行動"),
        ],
    },
    {
        "title": "情境五  接下來的收入可能變少",
        "situation": "你得知接下來一段時間的收入可能會變少，例如工時改變、轉換工作或收入不穩定。",
        "purpose": "觀察面對收入變動時，會重新計算、和重要的人協調，還是先維持現況。",
        "questions": [
            (13, "我會重新整理每月一定要付的錢和可以調整的花費，再安排接下來怎麼用。", "精算規劃"),
            (14, "我會和可能受到影響的人談談，看看生活要怎麼一起調整。", "關係照顧"),
            (15, "我可能先不整理哪些支出要調整，等影響真的出現後再處理。", "壓力迴避"),
        ],
    },
    {
        "title": "情境六  一起出錢的共同計畫",
        "situation": "你和家人或重要的人正在討論一個要一起花錢的計畫，例如旅行、搬家或添購家用品。",
        "purpose": "觀察共同做決定時，會先把握改變機會、照顧現在的生活，還是重視彼此都能接受。",
        "questions": [
            (16, "只要覺得這個計畫值得，我會願意先投入一小部分，再邊做邊調整。", "機會行動"),
            (17, "它能不能改善現在的生活，會影響我是否願意投入。", "當下感受"),
            (18, "我會先和對方談清楚預算與彼此期待，再一起決定要不要進行。", "關係照顧"),
        ],
    },
]


def set_font(run, name="Microsoft JhengHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=130, bottom=120, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_cell_width(cell, width):
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def style_table_text(cell, size=9.5, color=INK, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.25
        for run in paragraph.runs:
            set_font(run, size=size, bold=bold, color=color)


def add_header_row(table, labels, widths):
    row = table.rows[0]
    set_repeat_table_header(row)
    for index, (cell, label, width) in enumerate(zip(row.cells, labels, widths)):
        cell.text = label
        set_cell_width(cell, width)
        shade_cell(cell, PRIMARY)
        style_table_text(cell, size=9.5, color=WHITE, bold=True,
                         align=WD_ALIGN_PARAGRAPH.CENTER if index != 1 else WD_ALIGN_PARAGRAPH.LEFT)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    end_run = paragraph.add_run(" 頁")
    set_font(end_run, size=8.5, color=MUTED)


def add_paragraph(doc, text="", style=None, bold_lead=None, italic=False, align=None):
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_font(lead, size=11, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest, size=11, color=INK)
    else:
        run = paragraph.add_run(text)
        set_font(run, size=11, color=INK)
        run.italic = italic
    return paragraph


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_font(run, size=10.5, color=INK)
    return paragraph


def add_numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    run = paragraph.add_run(text)
    set_font(run, size=10.5, color=INK)
    return paragraph


def add_manual_numbered(doc, number, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.paragraph_format.first_line_indent = Cm(-0.55)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(f"{number}.  {text}")
    set_font(run, size=10.5, color=INK)
    return paragraph


def add_answer_options(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_with_next = True
    lead = paragraph.add_run("答案選項  ")
    set_font(lead, size=9.4, bold=True, color=PRIMARY)
    options = paragraph.add_run("／".join(ANSWER_OPTIONS))
    set_font(options, size=9.4, color=MUTED)
    return paragraph


def add_question_table(doc, questions):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(0.55), Inches(5.15), Inches(1.15)]
    add_header_row(table, ["題號", "題目", "觀察方向"], widths)
    for row_index, (number, question, direction) in enumerate(questions, start=1):
        row = table.add_row()
        prevent_row_split(row)
        values = [str(number), question, direction]
        for col_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.text = value
            set_cell_width(cell, width)
            if row_index % 2 == 0:
                shade_cell(cell, "FAF7F5")
            style_table_text(
                cell,
                size=9.5 if col_index == 1 else 9.2,
                color=INK,
                bold=col_index == 2,
                align=WD_ALIGN_PARAGRAPH.LEFT if col_index == 1 else WD_ALIGN_PARAGRAPH.CENTER,
            )
    set_table_borders(table)


def questions_for_plan(scenario, pressure_label):
    return [
        (number, question, pressure_label if direction == "壓力迴避" else direction)
        for number, question, direction in scenario["questions"]
    ]


def add_tool_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.15), Inches(2.15), Inches(3.55)]
    add_header_row(table, ["決策類型", "可以先使用", "另外兩項"], widths)
    rows = [
        ("安心守護", "財務健康儀表板", "儲蓄目標試算、財務韌性檢測"),
        ("精算規劃", "財務月報表", "生活目標財務規劃、儲蓄目標試算"),
        ("當下感受", "記帳助理", "財務月報表、生活目標財務規劃"),
        ("關係照顧", "財務月報表", "生活目標財務規劃、財務合拍檢測"),
        ("機會行動", "生活目標財務規劃", "時間資源盤點助理、財務健康儀表板"),
        ("壓力迴避", "財務健康儀表板", "財務月報表、時間資源盤點助理"),
    ]
    for row_index, row_values in enumerate(rows, start=1):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, (cell, value, width) in enumerate(zip(row.cells, row_values, widths)):
            cell.text = value
            set_cell_width(cell, width)
            if row_index % 2 == 0:
                shade_cell(cell, SAGE_LIGHT)
            style_table_text(cell, size=9.3, color=INK, bold=col_index == 0,
                             align=WD_ALIGN_PARAGRAPH.LEFT)
    set_table_borders(table)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True

    title_style = styles["Title"]
    title_style.font.name = "Microsoft JhengHei"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    title_style.paragraph_format.space_after = Pt(12)
    title_style_ppr = title_style._element.get_or_add_pPr()
    for border in title_style_ppr.findall(qn("w:pBdr")):
        title_style_ppr.remove(border)

    for style_name, size, before, after in (("Heading 1", 17, 18, 8), ("Heading 2", 13.5, 13, 5), ("Heading 3", 11.5, 10, 4)):
        style = styles[style_name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[list_style_name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_para.add_run("好理家在  財務決策檢測討論稿")
    set_font(header_run, size=8.5, color=MUTED)

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    properties = doc.core_properties
    properties.title = "財務決策檢測設計討論稿"
    properties.subject = "財務決策檢測題目與結果呈現確認"
    properties.author = "好理家在"
    properties.keywords = "財務決策, 檢測題目, 結果呈現"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("財務決策檢測設計討論稿")
    set_font(title_run, size=25, bold=True, color=BLACK)
    title_ppr = title._p.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    subtitle_run = subtitle.add_run("做金錢決定時 我會在意什麼  2026年9月17日")
    set_font(subtitle_run, size=11, color=MUTED)

    add_paragraph(
        doc,
        "這份討論稿整理六個生活情境、18題題目，以及完成作答後的結果呈現方式。希望先確認題目是否容易產生共鳴、分類名稱是否貼近原意，再進入頁面製作。",
    )
    add_paragraph(
        doc,
        "目前建議採用六種財務決策類型，包括安心守護、精算規劃、當下感受、關係照顧、機會行動及壓力迴避。壓力迴避描述的是壓力變大時可能先暫停或延後處理的反應，不代表消極或不負責。",
        bold_lead="目前建議",
    )

    doc.add_heading("一  檢測目的與作答方式", level=1)
    add_paragraph(doc, "檢測名稱暫定為「做金錢決定時，我會在意什麼？」。它不是理財能力測驗，也沒有標準答案，主要整理一個人在不同生活情境中，通常如何判斷、行動，以及壓力上來時可能怎麼回應。")
    for item in [
        "六個生活情境，共18題，每個情境3題。",
        "每個題組先交代一個生活情境，題目只問接下來會怎麼想或怎麼做，不重複同一句背景。正式作答時，當前情境會持續顯示，避免題目失去上下文。",
        "固定選項為很不像我、比較不像我、比較像我、很像我、我沒有遇過這個狀況。本報告在每個情境列一次，正式作答頁會隨當前題目顯示。",
        "結果只描述這次的作答，不顯示高低分，也不和其他人比較。",
        "重新檢測後從第一題開始，不保留前一次答案。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("二  題目設計原則", level=1)
    add_paragraph(doc, "以下完整列出六個情境、18題及五個答案選項。表格中的觀察方向只用來確認題目設計，不會顯示給填答者。")
    doc.add_heading("題目調整重點", level=2)
    for item in [
        "每個情境都有清楚的生活事件，題目只承接這筆錢、這個選擇或這個計畫，不重複整段背景。",
        "題目改用留下、分配、詢問、計算、延後、找方法及談預算等能想像的日常行動，減少只寫在意、照顧或踏出一步等抽象說法。",
        "情境三改為需要花錢且很快要決定的機會，避免把購物優惠和工作或課程機會混成同一種經驗。",
        "情境六明確寫成需要一起出錢的計畫，並提供旅行、搬家及家用品等例子，讓不同生活經驗的人比較容易代入。",
        "目前版本可作為題目確認稿；正式定稿前仍建議以少量試填確認是否有人對同一題產生不同理解。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("三  檢測架構  六種財務決策類型", level=1)
    add_paragraph(doc, "六種類型共同整理做金錢決定時通常會先在意什麼，以及壓力出現時可能採取的反應。每種類型都可能成為主要或次要結果。")
    doc.add_heading("完整題目", level=2)
    add_paragraph(doc, "第5、10、15題整理壓力迴避的傾向，和其他五個方向一起形成六種類型。")
    for scenario_index, scenario in enumerate(SCENARIOS):
        scenario_heading = doc.add_heading(scenario["title"], level=3)
        if scenario_index in (1, 3, 5):
            scenario_heading.paragraph_format.page_break_before = True
        add_paragraph(doc, f"生活情境  {scenario['situation']}", bold_lead="生活情境")
        add_paragraph(doc, f"設計目的  {scenario['purpose']}", italic=True)
        add_answer_options(doc)
        add_question_table(doc, questions_for_plan(scenario, "壓力迴避"))
    doc.add_heading("六種財務決策類型", level=2)
    for item in [
        "安心守護  先確認生活能不能穩穩過下去。",
        "精算規劃  先把數字、順序及後續影響看清楚。",
        "當下感受  先回應眼前真正需要被照顧的事情。",
        "關係照顧  把重要的人及共同生活放進考量。",
        "機會行動  看見可能帶來的改變，也願意找一個小步驟開始。",
        "壓力迴避  壓力變大時，可能先暫停或延後處理，等比較有空間時再面對。",
    ]:
        add_bullet(doc, item)
    doc.add_heading("結果頁呈現", level=2)
    for item in [
        "首屏從六種類型中整理主要與次要方向，例如「壓力迴避型 × 安心守護型」。",
        "雷達圖使用六個方向，呈現這次作答的描述性平均，不顯示高低分或與他人比較。",
        "每種類型都提供整體描述、帶來的力量、可以留意的地方、反思問題及一個小步驟。",
        "壓力迴避型的文字避免使用負面標籤，也不會自動導向焦慮、真人諮詢、急難或債務服務。",
        "工具依主要與次要類型提供起點，再由使用者確認目前最想處理的事情。",
    ]:
        add_bullet(doc, item)
    doc.add_heading("結果範例", level=2)
    plan_two_example = add_paragraph(doc, "壓力迴避型 × 安心守護型", bold_lead="壓力迴避型 × 安心守護型")
    plan_two_example.paragraph_format.keep_with_next = True
    add_paragraph(doc, "當金錢事情帶來壓力時，你可能會先讓自己喘口氣，等比較有空間時再處理；同時，你也很在意基本生活能不能穩穩維持。若事情有期限，可以先確認金額或截止日，其他部分再慢慢整理。")

    doc.add_heading("四  結果頁內容順序", level=1)
    steps = [
        "先呈現這次較接近的結果名稱與一句整體描述。",
        "呈現描述性雷達圖，不顯示高低分、百分位或他人比較。",
        "整理可能怎麼走到決定、這個模式帶來的力量及可以多留意的地方。",
        "提供一個反思問題與現在可以先做的小步驟。",
        "壓力迴避型和其他五種類型同樣可成為主要或次要結果，文案避免使用負面標籤。",
        "提供一項可以先開始的工具與兩項其他可用工具。",
        "最後才提供完整類型說明與重新檢測。",
    ]
    for step in steps:
        add_numbered(doc, step)

    doc.add_heading("五  工具推薦方式", level=1)
    add_paragraph(doc, "工具用來承接使用者目前的考量，不用最低分或弱項決定推薦。只有主要方向足夠清楚時，才使用類型對應；多種類型接近時，先回到一般整理工具。")
    add_tool_table(doc)
    add_paragraph(doc, "財務合拍檢測只在使用者確實有合適、也願意一起談錢的人時呈現。壓力迴避型的工具重點是先降低整理負擔，不自動導向焦慮、真人諮詢、急難或債務服務。")

    doc.add_heading("六  暫定整理原則", level=1)
    for item in [
        "六種類型各有3題，題目權重相同。",
        "四個相似程度選項可作為1到4的描述性數值，但前台不顯示數字。",
        "我沒有遇過這個狀況不納入平均，也不補成最低分。",
        "每種類型以有效題目的描述性平均整理，至少需要2題可用回答。",
        "某種類型少於2題可用回答時，標示資料不足，不補成0分或中間值；雷達圖如何呈現該軸仍需確認。",
        "一種類型明顯較高時，顯示主要類型；第二類型是否顯示，仍需訂出合理條件。",
        "兩種類型接近時可以並列；三種以上接近時，寫成「這次有多種類型接近」，不直接推論為靈活或適應力高。",
        "目前規則屬試行整理方式，正式上線前不宣稱為科學人格測驗或財務診斷。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("七  後續確認事項", level=1)
    add_paragraph(doc, "下列項目會影響正式名稱、結果文案與推薦規則，建議在進入頁面製作前確認。")
    decisions = [
        "這六個生活情境是否貼近實際服務對象？其中是否有一個情境較難產生共鳴？",
        "第5、10、15題是否有說中原本想表達的壓力迴避狀態？",
        "「壓力迴避型」作為第六種類型的名稱是否貼近原意？是否需要補上不等於消極或不負責的說明？",
        "「當下感受」目前較像在整理眼前生活照顧，「機會行動」較像先找一個做得到的方法開始；這兩個名稱是否需要調整？",
        "結果是否可以同時呈現兩個接近方向，或多種考量接近，而不是一定給出唯一類型？",
        "工具推薦應依主要類型提供起點，還是先讓使用者選擇目前最想處理的事情？",
    ]
    for number, item in enumerate(decisions, start=1):
        add_manual_numbered(doc, number, item)

    count_heading = doc.add_heading("八  題數建議", level=1)
    count_heading.paragraph_format.space_before = Pt(10)
    count_heading.paragraph_format.space_after = Pt(4)
    count_note = add_paragraph(doc, "建議先保留18題做小規模試填，記錄完成時間，以及看不懂或無法選擇的題目，再依回饋刪減。若縮成12題，每個方向只剩2題，只要一題回答「我沒有遇過這個狀況」，就難以整理。六種類型名稱及情境共鳴度確認前，先不定稿雷達圖、正式插圖或雙型組合文案。")
    count_note.paragraph_format.line_spacing = 1.25

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
